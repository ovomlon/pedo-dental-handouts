# -*- coding: utf-8 -*-
"""
生成《固定式间隙保持器与改良 Nance 弓矫治器——适应证·器械材料·操作步骤·禁忌证》
专业临床规范参考文档（DOCX），随后由 LibreOffice 转为 PDF。
"""
import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_EA = '宋体'
HEAD_EA = '微软雅黑'
BODY_ASCII = 'Times New Roman'
HEAD_ASCII = 'Arial'

DARK = RGBColor(0x1F, 0x36, 0x4D)
ACCENT = RGBColor(0x0F, 0x5C, 0x8A)
GREY = RGBColor(0x59, 0x59, 0x59)


# ---------------------------------------------------------------- helpers
def set_run(run, size=10.5, ea=BODY_EA, ascii_f=BODY_ASCII, bold=False,
            italic=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = ascii_f
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), ascii_f)
    rFonts.set(qn('w:hAnsi'), ascii_f)
    rFonts.set(qn('w:eastAsia'), ea)
    rFonts.set(qn('w:cs'), ascii_f)
    return run


def style_font(style, size, ea, ascii_f, bold=False, color=None):
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.name = ascii_f
    if color is not None:
        style.font.color.rgb = color
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), ascii_f)
    rFonts.set(qn('w:hAnsi'), ascii_f)
    rFonts.set(qn('w:eastAsia'), ea)
    rFonts.set(qn('w:cs'), ascii_f)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('start', left), ('bottom', bottom), ('end', right)):
        node = OxmlElement('w:' + tag)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tcPr.append(mar)


def add_field(paragraph, instr, size=9, ea=BODY_EA, ascii_f=BODY_ASCII):
    run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    set_run(run, size=size, ea=ea, ascii_f=ascii_f)
    return run


def make_styles(doc):
    st = doc.styles
    style_font(st['Normal'], 10.5, BODY_EA, BODY_ASCII)

    def mk(name, size, ea, ascii_f, bold, color, before, after, line=1.35,
           indent_chars=0, left=0.0, keep_next=False, base='Normal'):
        s = st.add_style(name, 1)  # paragraph style
        s.base_style = st[base]
        style_font(s, size, ea, ascii_f, bold, color)
        pf = s.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line
        if left:
            pf.left_indent = Cm(left)
        if indent_chars:
            pPr = s.element.get_or_add_pPr()
            ind = pPr.find(qn('w:ind'))
            if ind is None:
                ind = OxmlElement('w:ind')
                pPr.insert_element_before(
                    ind, 'w:contextualSpacing', 'w:mirrorIndents',
                    'w:suppressOverlap', 'w:jc', 'w:textDirection',
                    'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl',
                    'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
            ind.set(qn('w:firstLineChars'), str(int(indent_chars * 100)))
        if keep_next:
            s.paragraph_format.keep_with_next = True
        return s

    mk('CN-Title', 20, HEAD_EA, HEAD_ASCII, True, DARK, 0, 6, 1.15)
    mk('CN-Subtitle', 12.5, HEAD_EA, HEAD_ASCII, False, ACCENT, 0, 4, 1.2)
    mk('CN-H1', 15, HEAD_EA, HEAD_ASCII, True, DARK, 16, 8, 1.2, keep_next=True)
    mk('CN-H2', 12, HEAD_EA, HEAD_ASCII, True, ACCENT, 11, 5, 1.2, keep_next=True)
    mk('CN-H3', 11, HEAD_EA, HEAD_ASCII, True, RGBColor(0x33, 0x33, 0x33), 8, 4,
       1.2, keep_next=True)
    mk('CN-Body', 10.5, BODY_EA, BODY_ASCII, False, None, 0, 5, 1.4, indent_chars=2)
    mk('CN-List', 10.5, BODY_EA, BODY_ASCII, False, None, 0, 3, 1.35, left=0.75)
    mk('CN-Step', 10.5, BODY_EA, BODY_ASCII, False, None, 0, 3, 1.35, left=0.75)
    mk('CN-Caption', 9.5, HEAD_EA, HEAD_ASCII, True, DARK, 7, 3, 1.15,
       keep_next=True)
    mk('CN-Note', 9.5, BODY_EA, BODY_ASCII, False, GREY, 4, 6, 1.3, left=0.5)
    mk('CN-TblHead', 9, HEAD_EA, HEAD_ASCII, True, RGBColor(0xFF, 0xFF, 0xFF), 0, 0, 1.15)
    mk('CN-TblCell', 9, BODY_EA, BODY_ASCII, False, None, 0, 0, 1.2)
    mk('CN-Ref', 9.5, BODY_EA, BODY_ASCII, False, None, 0, 4, 1.3, left=0.75)


def title_block(doc, text, style, align=WD_ALIGN_PARAGRAPH.LEFT, after=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    set_run(p.add_run(text))
    # style carries the font; run needs eastAsia too
    for r in p.runs:
        r.font.size = Pt(doc.styles[style].font.size.pt)
        r.font.bold = doc.styles[style].font.bold
        rPr = r._element.get_or_add_rPr()
        rf = rPr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rPr.append(rf)
        rf.set(qn('w:eastAsia'), HEAD_EA if style != 'CN-Body' else BODY_EA)
    return p


def H1(doc, t):
    return title_block(doc, t, 'CN-H1')


def H2(doc, t):
    return title_block(doc, t, 'CN-H2')


def H3(doc, t):
    return title_block(doc, t, 'CN-H3')


def P(doc, t):
    p = doc.add_paragraph(style='CN-Body')
    set_run(p.add_run(t))
    return p


def B(doc, t, bold_prefix=None):
    p = doc.add_paragraph(style='CN-List')
    pf = p.paragraph_format
    pf.left_indent = Cm(1.1)
    pf.first_line_indent = Cm(-0.45)
    if bold_prefix:
        set_run(p.add_run(bold_prefix), bold=True)
    set_run(p.add_run(t))
    return p


def STEP(doc, num, t, bold_prefix=None):
    """编号步骤段落；bold_prefix 仅保留兼容签名，不再重复输出（避免与正文重复）。"""
    p = doc.add_paragraph(style='CN-Step')
    pf = p.paragraph_format
    pf.left_indent = Cm(1.25)
    pf.first_line_indent = Cm(-0.75)
    set_run(p.add_run(num + '　'), bold=True, color=ACCENT)
    set_run(p.add_run(t))
    return p


def NOTE(doc, t, prefix='注意：'):
    p = doc.add_paragraph(style='CN-Note')
    set_run(p.add_run(prefix), bold=True, color=RGBColor(0xB0, 0x3A, 0x2B))
    set_run(p.add_run(t))
    return p


def CAP(doc, t):
    return title_block(doc, t, 'CN-Caption')


def TABLE(doc, cols, widths, rows, first_col_bold=False, size=9):
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for i, c in enumerate(cols):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(c), size=size, ea=HEAD_EA, ascii_f=HEAD_ASCII,
                bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(hdr[i], '1F364D')
        cell_margins(hdr[i])
    for r in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(r):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            bold = first_col_bold and i == 0
            set_run(p.add_run(txt), size=size, bold=bold)
            cell_margins(cells[i])
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    # repeat header row on page breaks
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    trPr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- document
doc = Document()
make_styles(doc)

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)
sec.header_distance = Cm(1.2)
sec.footer_distance = Cm(1.0)

# header
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(hp.add_run('固定式间隙保持器与改良 Nance 弓矫治器 · 儿童口腔临床操作规范参考'),
        size=8.5, ea=HEAD_EA, ascii_f=HEAD_ASCII, color=GREY)
pPr = hp._p.get_or_add_pPr()
pbdr = OxmlElement('w:pBdr')
bot = OxmlElement('w:bottom')
bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '4')
bot.set(qn('w:space'), '2'); bot.set(qn('w:color'), 'BFBFBF')
pbdr.append(bot)
pPr.insert_element_before(pbdr, 'w:shd', 'w:tabs', 'w:suppressAutoHyphens',
                          'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
                          'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
                          'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
                          'w:spacing', 'w:ind', 'w:contextualSpacing',
                          'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
                          'w:textDirection', 'w:textAlignment',
                          'w:textboxTightWrap', 'w:outlineLvl', 'w:divId',
                          'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')

# footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(fp.add_run('第 '), size=9)
add_field(fp, 'PAGE')
set_run(fp.add_run(' 页 / 共 '), size=9)
add_field(fp, 'NUMPAGES')
set_run(fp.add_run(' 页　|　仅供临床参考，具体病例请遵循最新指南与个体化评估'), size=8.5,
        color=GREY)

TODAY = datetime.date.today().strftime('%Y-%m-%d')

# ============================ 封面区 ============================
title_block(doc, '固定式间隙保持器与改良 Nance 弓矫治器', 'CN-Title',
            WD_ALIGN_PARAGRAPH.CENTER)
title_block(doc, '适应证 · 器械材料 · 操作步骤 · 禁忌证', 'CN-Subtitle',
            WD_ALIGN_PARAGRAPH.CENTER)
title_block(doc, '——儿童口腔临床操作规范参考（专业版）', 'CN-Subtitle',
            WD_ALIGN_PARAGRAPH.CENTER)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

TABLE(doc,
      ['项目', '内容'],
      [3.2, 13.3],
      [['文件性质', '临床操作规范参考 / 教学资料（非替代性诊疗文件）'],
       ['适用对象', '儿童口腔科医师、口腔全科医师、口腔正畸医师、口腔医学生及护理人员'],
       ['适用人群', '乳牙列期与混合牙列期的间隙管理、早期阻断性矫治患者'],
       ['编制依据', '国内外间隙管理指南与专家共识、儿童口腔医学与口腔正畸学教材、'
                    'AAPD 儿童牙列与咬合发育管理相关规范'],
       ['核心内容', '两类装置的适应证、禁忌证、器械材料清单、临床与技工室操作步骤、'
                    '复诊维护与并发症处理'],
       ['版本日期', TODAY + '（第 1 版）']],
      first_col_bold=True, size=9.5)

NOTE(doc, '本文件所列药物、材料规格、力值与复诊间隔均为临床常用参考范围，'
          '实际应用须结合患者年龄、牙列阶段、影像学表现、口腔卫生状况与依从性个体化调整；'
          '涉及具体产品时以厂家说明书与注册证载明内容为准。', prefix='使用说明：')

# ============================ 目录导览 ============================
H1(doc, '内容导览')
B(doc, '第 1 章　概述与装置分类　|　第 2 章　适应证与装置选择　|　第 3 章　禁忌证')
B(doc, '第 4 章　器械与材料　|　第 5 章　操作步骤（临床 + 技工室）')
B(doc, '第 6 章　复诊、维护与并发症处理　|　第 7 章　拆除指征与时机')
B(doc, '第 8 章　质量控制要点　|　第 9 章　附录（核对清单、材料速查、家长告知）　|　参考文献')

# ============================ 1 概述 ============================
H1(doc, '1　概述与装置分类')

H2(doc, '1.1　间隙管理的基本原理')
P(doc, '乳牙在儿童牙列中承担咀嚼、发音、引导恒牙萌出与维持牙弓长度等多重功能。'
       '乳牙因龋病、牙髓根尖周病、外伤或全身因素过早丧失后，缺隙两侧邻牙可向缺隙倾斜、'
       '对颌牙可过度萌出、继承恒牙可因间隙不足而阻生或异位萌出，最终导致牙弓长度缩短、'
       '咬合关系紊乱与中线偏移。间隙保持的核心目标是：在继承恒牙萌出或综合正畸治疗介入之前，'
       '维持缺隙的近远中径与垂直向高度，防止间隙丧失与继发错合畸形。')
P(doc, '间隙丧失的速度与程度受多因素影响：乳磨牙早失后最初 6 个月间隙丧失最为迅速；'
       '第一恒磨牙萌出期（"第一恒磨牙推挤期"）间隙关闭风险最高；'
       '第二乳磨牙早失、缺隙远中无牙阻挡、对颌牙过度萌出、口腔不良习惯与龋坏未控制者，'
       '间隙丧失风险进一步升高。因此，间隙保持的决策应有时间敏感性，宜在乳牙拔除或早失后尽早完成。')

H2(doc, '1.2　固定式与活动式间隙保持器的比较')
B(doc, '固定式：以带环、全冠或粘接方式固定于基牙，患者不能自行摘戴；'
       '依从性依赖度低，保持效果相对稳定，适用于依从性差、需长期保持、多牙缺失或跨中线缺隙；'
       '缺点为清洁困难、龋与牙龈炎风险增加、需复诊维护、装置脱落时不易被家长察觉。')
B(doc, '活动式：可自行摘戴，便于清洁与调整，适用于个别牙缺失且依从性良好者；'
       '缺点为易丢失、易损坏、依赖佩戴时间，依从性差时保持失败率高。')

H2(doc, '1.3　固定式间隙保持器的主要类型')
CAP(doc, '表 1　固定式间隙保持器的主要类型与适用范围')
TABLE(doc,
      ['类型', '主要适用情形', '常用支抗牙', '跨中线'],
      [2.9, 8.0, 4.1, 1.5],
      [['带环丝圈式\n（Band and loop）',
        '单侧 1 个乳磨牙早失，继承恒牙预计 6 个月以上不能萌出；基牙临床牙冠高度与形态良好',
        '缺隙远中邻牙（乳磨牙或恒磨牙）', '否'],
       ['全冠丝圈式',
        '同带环丝圈式，但基牙大面积缺损、已行预成冠或需全冠修复者',
        '全冠修复后的基牙', '否'],
       ['远端靴式\n（Distal shoe）',
        '第一恒磨牙萌出前，第二乳磨牙早失（下颌尤为经典）；需引导第一恒磨牙正常萌出',
        '第一乳磨牙（必要时延伸至龈下）', '否'],
       ['固定舌弓\n（Lingual arch）',
        '双侧下颌乳磨牙早失、多牙缺失；需保持下颌牙弓长度与磨牙间宽度',
        '双侧下颌第一恒磨牙（或第二乳磨牙）', '是'],
       ['Nance 腭弓',
        '双侧上颌乳磨牙早失；需跨中线保持上颌牙弓长度；腭穹窿形态利于腭托就位',
        '双侧上颌第一恒磨牙', '是'],
       ['改良 Nance 弓矫治器',
        '间隙保持合并早期阻断或支抗需求：磨牙异位萌出、推磨牙向远中、前方牵引支抗、'
        '扩弓后保持等（详见 2.3）',
        '双侧上颌第一恒磨牙，必要时加腭部骨支抗', '是'],
       ['横腭杆（TPA）',
        '维持上颌磨牙间宽度、加强磨牙支抗，多与固定矫治联合应用',
        '双侧上颌磨牙', '是'],
       ['组合式装置',
        '复杂缺隙或需同时实现扩弓、导板、习惯阻断等功能时的个性化设计',
        '依设计而定', '依设计']],
      first_col_bold=True)

NOTE(doc, '本文件重点阐述①固定式间隙保持器（以带环丝圈式、固定舌弓、Nance 腭弓为代表）'
          '与②改良 Nance 弓矫治器两类装置；远端靴式、横腭杆等仅作分类学对照。')

# ============================ 2 适应证 ============================
H1(doc, '2　适应证与装置选择')

H2(doc, '2.1　固定式间隙保持器的通用适应证')
B(doc, '单侧或双侧 1 个及以上乳磨牙早失，影像学显示继承恒牙距萌出仍有较长时间'
       '（一般要求牙根形成不足 1/2～2/3，或恒牙牙冠上方仍有骨质覆盖、未突破牙龈）。')
B(doc, '预计缺隙在继承恒牙萌出前会明显缩小：如第二乳磨牙早失而第一恒磨牙正在萌出、'
       '缺隙远中缺少牙体阻挡、对颌牙伸长。')
B(doc, '乳牙早失同时伴有以下高危因素：口腔不良习惯、肌功能异常、邻牙已出现倾斜趋势、'
       '牙弓长度已轻度缩短需阻止其进一步发展。')
B(doc, '先天缺牙（如先天缺失前磨牙、侧切牙）需维持间隙至正畸方案确定或成年后修复/种植。')
B(doc, '多牙缺失、跨中线缺隙，或儿童年龄小、依从性不足，无法有效使用活动式保持器。')
B(doc, '作为早期阻断性矫治的组成部分：如需要加强磨牙支抗以控制磨牙近中移动、'
       '配合前方牵引或扩弓治疗。')

H2(doc, '2.2　各类型装置的适应证细则')
H3(doc, '2.2.1　带环丝圈式（含全冠丝圈式）')
B(doc, '单侧一个牙位缺失，通常为第一乳磨牙或第二乳磨牙早失；')
B(doc, '基牙为缺隙远中的乳磨牙或恒磨牙，牙冠完整、高度足够、无大面积充填体或已行全冠修复；')
B(doc, '继承恒牙在影像学上距萌出尚有 6 个月以上，且无阻生、异位萌出或需早期拔除的指征；')
B(doc, '缺隙近中邻牙存在且牙体健康，丝圈可稳定抵靠于其邻面。')
H3(doc, '2.2.2　固定舌弓')
B(doc, '双侧下颌乳磨牙早失（最常见为双侧第二乳磨牙早失），需同时保持双侧间隙与牙弓长度；')
B(doc, '下颌多个牙位缺失、下颌牙弓长度有缩短趋势，或需要保持替牙间隙（Leeway space）；')
B(doc, '下颌第一恒磨牙已萌出且牙冠条件满足带环固位要求；')
B(doc, '需避免使用活动式装置（依从性差、年龄小、易丢失）。')
H3(doc, '2.2.3　Nance 腭弓')
B(doc, '双侧上颌乳磨牙早失、需跨中线同时保持双侧缺隙与上颌牙弓长度；')
B(doc, '上颌第一恒磨牙已萌出或萌出程度足以支持带环固位；')
B(doc, '腭穹窿有一定深度、腭黏膜健康，可容纳前腭托而不引起压迫或发音障碍；')
B(doc, '需防止上颌磨牙近中移动、维持磨牙关系者。')
H3(doc, '2.2.4　远端靴式（分类学对照）')
B(doc, '下颌第二乳磨牙早失而第一恒磨牙尚未萌出，需以金属延伸臂引导第一恒磨牙在正常位置萌出；')
B(doc, '经典术式要求在拔牙前取模、拔牙后即刻戴入，对时机与技工配合要求高。')

H2(doc, '2.3　改良 Nance 弓矫治器的适应证（重点）')
P(doc, '改良 Nance 弓矫治器是在经典 Nance 腭弓（双侧上颌第一恒磨牙带环 + 腭弓 + 前腭托）'
       '基础上，通过增加功能部件或改变结构与固位方式而制成的固定式保持/矫治一体装置。'
       '其核心优势在于：腭托与腭弓提供稳定支抗，同时通过改良部件实现间隙保持之外的治疗目标。')
B(doc, '上颌双侧乳磨牙早失的间隙保持，尤其腭穹窿较高、需要可靠跨中线支抗、'
       '且家长难以保证活动式装置佩戴时间者；')
B(doc, '上颌第一恒磨牙异位萌出（ectopic eruption）的早期阻断：于腭弓或带环上焊接推簧、'
       '圈曲或远中向牵引臂，对异位萌出的第一恒磨牙施加可控的远中向力，'
       '同时腭托提供支抗、防止支抗牙近中移动；')
B(doc, '早期安氏 II 类（上颌磨牙近中移位、磨牙远中关系）的推磨牙向远中治疗，'
       '改善磨牙关系，为后续治疗创造条件；')
B(doc, '安氏 III 类上颌发育不足（上颌后缩）行前方牵引（面罩）治疗时的腭部支抗：'
       '于带环上焊接颊面管或牵引钩，配合面具施加前牵引力，腭托分散反作用力、'
       '减少磨牙近中移动与牙弓长度丧失；')
B(doc, '上颌横向发育不足且合并间隙丧失时，与螺旋扩弓器结合形成"扩弓—保持"一体装置；')
B(doc, '后牙反合、锁合早期阻断及扩弓治疗后的效果维持；')
B(doc, '基牙大面积缺损、单侧缺隙或牙冠形态异常时的固位改良设计（改全冠、增加邻牙支托、'
       '半侧/不对称设计）；')
B(doc, '需腭部阻挡以协助纠正异常吞咽、吐舌等口腔不良习惯，或需附加平面导板改善深覆合者。')

H2(doc, '2.4　装置选择决策速查')
CAP(doc, '表 2　常见临床情形与装置选择建议')
TABLE(doc,
      ['临床情形', '首选考虑', '选择要点'],
      [5.6, 4.4, 6.5],
      [['单侧下颌第一乳磨牙早失，第二乳磨牙健康',
        '带环丝圈式', '支抗牙为第二乳磨牙；丝圈跨过缺隙抵于近中邻牙邻面'],
       ['单侧下颌第二乳磨牙早失，第一恒磨牙未萌',
        '远端靴式', '需引导第一恒磨牙萌出；对时机与技工配合要求高'],
       ['双侧下颌乳磨牙早失',
        '固定舌弓', '支抗牙为双侧下颌第一恒磨牙；注意舌侧软组织避让'],
       ['双侧上颌乳磨牙早失',
        'Nance 腭弓', '腭托需与腭黏膜均匀接触、边缘圆钝'],
       ['上颌第一恒磨牙异位萌出（可复性/嵌顿）',
        '改良 Nance 弓（加推簧或牵引臂）', '宜早期干预；需评估磨牙可否被推回'],
       ['上颌后缩（安氏 III 类）拟行前方牵引',
        '改良 Nance 弓（加颊面管/牵引钩）+ 面具', '腭托提供支抗；需保证每日佩戴时间'],
       ['上颌横向不足合并间隙丧失',
        '改良 Nance 弓 + 螺旋扩弓器', '扩弓与保持同步，注意加力节奏与腭中缝反应'],
       ['口腔卫生极差、龋病未控制',
        '暂缓固定装置，先控制龋病', '改用活动式或延长观察期，强化菌斑控制']],
      first_col_bold=True)

# ============================ 3 禁忌证 ============================
H1(doc, '3　禁忌证')

H2(doc, '3.1　通用禁忌证（绝对或需先行处理）')
B(doc, '继承恒牙即将萌出：影像学显示恒牙牙根形成超过 2/3、骨覆盖已被突破、'
       '或临床可触及萌出中的牙尖，预计 6 个月内萌出者，装置将阻碍其正常萌出。')
B(doc, '未控制的猖獗龋、急性牙髓炎或根尖周炎、口腔急性炎症：须先完成龋病与牙髓治疗、'
       '控制感染后再行修复性/保持性装置。')
B(doc, '严重的全身性疾病、免疫缺陷、出血性疾病未控制，或无法耐受长时间张口操作者。')
B(doc, '对镍铬合金、不锈钢等装置材料明确过敏者（可考虑钛合金、贵金属或全瓷全冠型设计，'
       '属相对禁忌）。')
B(doc, '患者或监护人拒绝治疗、无法理解并配合维护要求，或明确无定期复诊条件。')

H2(doc, '3.2　相对禁忌证与慎用情形')
B(doc, '口腔卫生差、菌斑指数高、高龋活性：固定装置会加剧菌斑堆积与带环下脱矿，'
       '应先进行口腔卫生指导、洁治与氟化物处理，并优先选择易于清洁的设计，'
       '必要时暂缓或改用活动式装置。')
B(doc, '基牙条件不足：临床牙冠过短、锥形牙、大面积缺损未修复、严重磨耗，'
       '无法获得带环/全冠的有效固位。')
B(doc, '第一恒磨牙尚未萌出或萌出不足而计划以带环固位者（应改用远端靴式或其他设计）。')
B(doc, '腭穹窿过浅、腭部软组织增生、腭裂修复术后瘢痕或腭部感觉/吞咽功能障碍者，'
       'Nance 腭托易压迫、脱位或引起不适。')
B(doc, '严重磨牙症、咬合力量异常或存在对合干扰，装置脱落率与变形率明显升高。')
B(doc, '混合牙列快速生长期、预计短期内需更换装置者，宜缩短复诊间隔并预告重做可能。')
B(doc, '已明确需要综合正畸治疗（含拔牙矫治）者：固定式间隙保持器可能干扰矫治计划，'
       '应与正畸医师共同决策，或直接进入综合矫治。')
B(doc, '存在需要优先处理的颌骨发育异常、多生牙、囊肿、牙瘤等病变者，应先处理病因。')

H2(doc, '3.3　改良 Nance 弓矫治器的特有禁忌与慎用')
B(doc, '腭穹窿过浅、腭托无法就位或反复脱位者，装置固位与支抗效果不可靠。')
B(doc, '腭部黏膜溃疡、腭部术后缺损、腭裂修复后瘢痕组织、腭部感觉或吞咽功能障碍'
       '（如脑性瘫痪）者，腭托可致压迫性溃疡、误吸风险增加。')
B(doc, '需前方牵引但无法保证面具每日有效佩戴时间（一般需 12～14 小时/日）者，'
       '治疗效果不可预期，应重新评估方案。')
B(doc, '后牙区缺乏有效支抗条件：上颌第一恒磨牙未萌、牙根发育不足、'
       '牙周支持组织不足或牙体大面积缺损。')
B(doc, '口腔卫生极差者：腭托组织面与带环周围是菌斑滞留区，'
       '易发生牙龈炎、腭部黏膜炎症与牙体脱矿。')
B(doc, '存在需外科或综合正畸处理的严重骨性错合（如明显上颌后缩伴下颌前突、'
       '骨性开合），单纯早期推磨牙或前牵引获益有限。')
B(doc, '监护人及患儿对持续加力可能带来的磨牙牙根吸收、牙髓反应、带环松动、'
       '脱矿等风险缺乏认识或不能配合随访者。')

# ============================ 4 器械与材料 ============================
page_break(doc)
H1(doc, '4　器械与材料')

H2(doc, '4.1　检查、取模与模型材料')
B(doc, '一次性口腔检查盘：口镜、探针、镊子；间隙测量尺或游标卡尺。')
B(doc, '儿童型号全口/局部托盘、藻酸盐印模材（或硅橡胶印模材）、调拌碗与调刀、'
       '或自动调拌机。')
B(doc, '超硬石膏（人造石）、模型修整机、模型锯、分离剂。')
B(doc, '咬合记录蜡、比色板（全冠/树脂修复时使用）。')

H2(doc, '4.2　固位体（带环与全冠）及粘接材料')
B(doc, '预成不锈钢带环（厚度常用 0.12～0.18 mm）、带环推压器、带环扩大钳、'
       '去冠器、咬合纸。')
B(doc, '预成不锈钢全冠（乳磨牙）、冠剪、锤冠器或成形钳。')
B(doc, '粘接材料：玻璃离子水门汀、树脂加强型玻璃离子水门汀、'
       '聚羧酸锌水门汀；必要时正畸粘接剂配合酸蚀技术。')
B(doc, '隔湿与防护：开口器、吸唾器、棉卷、橡皮障（条件允许时）、'
       '抛光杯与抛光膏、氟保护漆或含氟涂料。')

H2(doc, '4.3　技工室器械与材料')
B(doc, '不锈钢丝：0.8 mm（0.032 in）、0.9 mm（0.036 in）、1.0 mm；'
       '丝圈与舌弓常用 0.8～0.9 mm，腭弓常用 0.9～1.0 mm。')
B(doc, '弯制与切断工具：梯形钳、三喙钳、弯丝钳、断丝钳、平口钳。')
B(doc, '焊接设备：点焊机（电焊）、银焊或微焊焊料、焊媒（硼砂）、'
       '酒精灯或微型焊枪。')
B(doc, '腭托材料：自凝或热凝丙烯酸树脂、光固化复合树脂（对丙烯酸类敏感者可选）；'
       '蜡、雕刀、分离剂。')
B(doc, '打磨抛光：打磨机、砂石、橡皮轮、抛光膏、电解抛光设备。')

H2(doc, '4.4　改良 Nance 弓矫治器的附加部件')
B(doc, '螺旋扩弓器（如 Hyrax、Haas 型）：用于扩弓—保持一体设计；'
       '需在模型上预排位置后再包埋固定。')
B(doc, '颊面管、牵引钩、舌侧扣：用于前方牵引或牵引附件。')
B(doc, '推簧、镍钛丝、圈曲（loop）：用于推磨牙向远中或异位磨牙的复位。')
B(doc, '前方牵引面具（面罩）与弹力牵引橡皮圈（常用规格 3/8 in、5/16 in 等）。')
B(doc, '腭部微种植钉（可选）：需要绝对支抗或减少磨牙近中移动时使用。')

CAP(doc, '表 3　常用器械材料速查表')
TABLE(doc,
      ['类别', '名称与常用规格', '用途与要点'],
      [2.8, 6.3, 7.4],
      [['检查与取模', '检查盘、间隙测量尺、儿童托盘、藻酸盐印模材、超硬石膏',
        '带环就位后取模，带环须牢固嵌于印模内并准确复位'],
       ['固位体', '预成不锈钢带环（0.12～0.18 mm）、预成不锈钢全冠',
        '基牙固位；边缘密合、无悬突、无咬合干扰'],
       ['不锈钢丝', '0.8 mm / 0.9 mm / 1.0 mm',
        '丝圈、舌弓、腭弓弯制；直径越大刚度越高、调整越困难'],
       ['焊接', '点焊机、银焊/微焊、焊媒',
        '丝与带环连接；避免过热导致带环变形或退火'],
       ['腭托材料', '自凝/热凝丙烯酸树脂、复合树脂',
        '前腭托制作；厚度约 2～3 mm，边缘圆钝、均匀贴合'],
       ['粘接', '玻璃离子、树脂加强玻璃离子、正畸粘接剂',
        '带环粘固；隔湿是成功关键'],
       ['改良部件', '螺旋扩弓器、颊面管、牵引钩、推簧、面具',
        '实现推磨牙、前方牵引、扩弓等功能性改良'],
       ['防护', '开口器、吸唾器、棉卷、氟保护漆',
        '隔湿、防龋与术后脱矿预防']],
      first_col_bold=True)

# ============================ 5 操作步骤 ============================
page_break(doc)
H1(doc, '5　操作步骤')

H2(doc, '5.1　术前评估与准备（适用于所有固定式间隙保持器）')
STEP(doc, '第 1 步', '采集病史与临床检查：明确早失牙位置、原因（龋病、牙髓根尖周病、外伤、'
     '拔除）与时间；评估口腔卫生与龋活性、口腔不良习惯、咬合关系、中线、'
     '覆合覆盖、基牙牙冠条件及患儿配合度；询问全身病史与过敏史。')
STEP(doc, '第 2 步', '影像学评估：拍摄根尖片或全口曲面体层片（必要时锥形束 CT）；'
     '重点评估继承恒牙的发育阶段（牙根形成程度）、位置与萌出通道、骨覆盖情况、'
     '有无多生牙、先天缺牙、牙瘤、囊肿及异位萌出。')
STEP(doc, '第 3 步', '间隙与咬合记录：测量缺隙近远中径，记录牙弓长度、磨牙关系、'
     '覆合覆盖与中线偏移；必要时取诊断模型与咬合记录。')
STEP(doc, '第 4 步', '方案选择与知情同意：向监护人说明装置目的、类型选择理由、'
     '维护要求、脱落与并发症风险、复诊计划与费用；签署知情同意书。')
STEP(doc, '第 5 步', '口腔预处理：完成全口洁治与氟化物处理，治疗未控制的龋病与牙髓根尖周病；'
     '进行针对性菌斑控制指导与饮食指导。')

NOTE(doc, '术前应特别确认：继承恒牙距萌出仍有足够时间、基牙可提供有效固位、'
          '监护人能配合复诊与日常维护。三者之一不满足时，应重新评估方案。')

H2(doc, '5.2　带环丝圈式固定间隙保持器')
STEP(doc, '第 1 步', '选择并试戴带环：以缺隙远中的乳磨牙或恒磨牙为基牙，'
     '选取合适型号预成带环，试戴至完全就位；要求边缘密合无悬突、'
     '邻接关系恢复、咬合无干扰、不影响牙龈健康。', '试带环：')
STEP(doc, '第 2 步', '取印模：带环就位状态下取上下颌印模，取出印模后确认带环牢固嵌于印模内；'
     '若带环脱落或移位须重新取模。', '取模：')
STEP(doc, '第 3 步', '灌制模型：以超硬石膏灌注，带环准确复位固定于模型上，'
     '待石膏完全凝固后修整模型。', '灌模：')
STEP(doc, '第 4 步', '技工室弯制与焊接：以 0.8～0.9 mm 不锈钢丝弯制丝圈，'
     '一端焊接于带环颊面或舌面（焊点覆盖带环宽度 1/2 以上）；'
     '丝圈跨过缺隙，近中端形成抵靠于缺隙近中邻牙邻面的 U 形曲或环形结构。'
     '要求：丝圈距牙龈 0.5～1.0 mm、不进入邻牙倒凹、不压迫牙龈、'
     '不妨碍咬合、不形成菌斑滞留死角，并预留少量调整余地。', '弯制丝圈：')
STEP(doc, '第 5 步', '试戴与调整：口内试戴，检查就位、稳定性、邻接、咬合与清洁通道，'
     '必要时调整丝圈形态与位置。', '试戴：')
STEP(doc, '第 6 步', '粘接：隔湿、清洁并干燥基牙，以玻璃离子水门汀或树脂加强玻璃离子'
     '粘固带环，去除多余粘接剂，检查咬合与边缘密合度。', '粘接带环：')
STEP(doc, '第 7 步', '医嘱与预约：向监护人演示清洁方法（重点为带环边缘与丝圈下方），'
     '交代饮食注意事项与脱落时的处理，预约 1 周后复查。', '医嘱：')

H2(doc, '5.3　固定舌弓')
STEP(doc, '第 1 步', '试戴双侧带环：以下颌第一恒磨牙（或第二乳磨牙）为基牙，'
     '分别试戴带环并确认就位与咬合。', '试带环：')
STEP(doc, '第 2 步', '取模、灌模：同 5.2 第 2、3 步，双侧带环均需牢固复位。', '取模灌模：')
STEP(doc, '第 3 步', '技工室弯制：以 0.9 mm 不锈钢丝弯制舌弓，'
     '沿下颌舌侧牙冠外形走行，距牙龈与舌侧黏膜 0.5～1.0 mm；'
     '在中段设置 U 形曲以提供调整余地；'
     '舌弓与前牙舌面的关系按设计确定（接触式用于控制舌倾，非接触式用于单纯保持间隙）。',
     '弯制舌弓：')
STEP(doc, '第 4 步', '焊接：将舌弓两端分别焊接于双侧带环舌面，焊点牢固、无变形。', '焊接：')
STEP(doc, '第 5 步', '试戴、粘接与医嘱：口内试戴检查组织避让与咬合，'
     '隔湿后粘固，交代清洁与复诊要求。', '试戴与粘接：')

NOTE(doc, '舌弓中段的 U 形曲是后续调整间隙与适应牙弓生长的关键结构，'
          '制作时应避免曲部压迫舌侧牙龈或刺激舌体。')

H2(doc, '5.4　经典 Nance 腭弓')
STEP(doc, '第 1 步', '试戴双侧上颌带环：以上颌第一恒磨牙为基牙，试戴并确认就位、'
     '邻接与咬合。', '试带环：')
STEP(doc, '第 2 步', '取模、灌模：同 5.2 相应步骤，带环复位固定。', '取模灌模：')
STEP(doc, '第 3 步', '技工室弯制腭弓：以 0.9～1.0 mm 不锈钢丝沿腭穹窿弯制，'
     '距腭黏膜 0.5～1.0 mm，避免压迫腭皱襞与腭大神经血管束走行区；'
     '前端止于腭皱襞后方（大致相当于第二前磨牙区）。', '弯制腭弓：')
STEP(doc, '第 4 步', '制作前腭托：于模型腭前部放置自凝树脂，覆盖腭弓末端，'
     '形成厚度约 2～3 mm、边缘圆钝、表面光滑的腭托；'
     '腭托组织面须与腭黏膜均匀贴合，边缘不压迫牙龈、不越过后牙功能尖区、'
     '不影响舌体活动与发音。', '制作腭托：')
STEP(doc, '第 5 步', '焊接与抛光：将腭弓两端焊接于带环舌面，'
     '整体打磨抛光、电解抛光，避免粗糙面滞留菌斑。', '焊接抛光：')
STEP(doc, '第 6 步', '试戴：口内试戴，检查腭托贴合度与稳定性、咬合、发音与吞咽；'
     '必要时缓冲组织面。', '试戴：')
STEP(doc, '第 7 步', '粘接与医嘱：隔湿后粘固带环，去除多余粘接剂；'
     '交代清洁重点（腭托边缘与带环周围）与复诊计划。', '粘接：')

H2(doc, '5.5　改良 Nance 弓矫治器')

H3(doc, '5.5.1　常见改良形式与设计要点')
CAP(doc, '表 4　改良 Nance 弓矫治器的常见改良形式')
TABLE(doc,
      ['改良形式', '结构变化', '主要目的与适应证'],
      [3.0, 6.0, 7.5],
      [['推磨牙向远中型', '腭弓加圈曲或推簧，或焊接远中向牵引臂、舌侧扣',
        '上颌第一恒磨牙异位萌出、早期安氏 II 类磨牙关系的阻断性治疗'],
       ['支抗增强型（前牵引）', '带环焊接颊面管或牵引钩，腭托适当加大以分散反作用力',
        '安氏 III 类上颌后缩行前方牵引时的腭部支抗'],
       ['扩弓—保持型', '腭中缝区加装螺旋扩弓器（Hyrax/Haas 型）',
        '上颌横向发育不足合并间隙丧失，扩弓与间隙保持同步'],
       ['固位改良型', '基牙改为全冠、增加邻牙支托或采用半侧/不对称设计',
        '基牙大面积缺损、单侧缺隙、牙冠形态异常'],
       ['骨支抗型', '腭部微种植钉或腭中缝种植钉辅助',
        '需要绝对支抗、控制磨牙近中移动的病例'],
       ['软组织/习惯改良型', '腭托延伸为腭屏或附加平面导板',
        '异常吞咽、吐舌习惯、深覆合的辅助纠正']],
      first_col_bold=True)

H3(doc, '5.5.2　技工室制作步骤')
STEP(doc, '第 1 步', '模型准备与设计标记：修整模型，标记带环位置、'
     '功能部件位置（推簧、牵引钩、扩弓器）与腭托范围；'
     '必要时预留扩弓器的就位空间。', '模型准备：')
STEP(doc, '第 2 步', '弯制腭弓并与带环焊接：先完成腭弓形态（同 5.4 第 3 步），'
     '再与带环舌面焊接；焊接时控制加热范围，防止带环变形退火。', '腭弓焊接：')
STEP(doc, '第 3 步', '制作腭托：同 5.4 第 4 步；注意避开腭中缝（未扩弓时）、'
     '腭大孔区域与腭皱襞，腭托不宜过度延伸以免影响舌体运动。', '制作腭托：')
STEP(doc, '第 4 步', '焊接功能部件：依次焊接推簧/圈曲、颊面管、牵引钩、扩弓器等；'
     '焊接顺序宜先大件后小件，焊后逐一检查位置、方向与强度；'
     '扩弓器须与腭中缝平行、与牙弓中线对称。', '焊接功能部件：')
STEP(doc, '第 5 步', '打磨、抛光与模型试合：电解抛光，'
     '在模型上试合上下颌检查咬合关系与部件是否干扰。', '打磨抛光：')

H3(doc, '5.5.3　临床操作步骤')
STEP(doc, '第 1 步', '术前评估：同 5.1；重点评估上颌第一恒磨牙的萌出程度与牙根发育、'
     '异位萌出的可复性、腭部软组织条件与患儿对前牵引/扩弓治疗的配合能力。', '术前评估：')
STEP(doc, '第 2 步', '试戴部件与取模：试戴带环及功能部件（颊面管、牵引钩等），'
     '确认位置后取模、灌模并交付技工室制作。', '试戴与取模：')
STEP(doc, '第 3 步', '口内试戴：检查带环就位、腭托贴合度、'
     '功能部件的位置与力线方向（推磨牙方向、前牵引角度）、咬合与发音；'
     '必要时缓冲或调整。', '试戴：')
STEP(doc, '第 4 步', '粘接：隔湿后以玻璃离子水门汀或树脂加强玻璃离子粘固带环；'
     '如需粘接颊面管或附件，按正畸粘接流程操作。', '粘接：')
STEP(doc, '第 5 步', '加力与首次调整：按设计加力。'
     '推磨牙向远中者，通常每次加力约 0.5～1.0 mm 或按装置设计施加轻力；'
     '前方牵引者，一般施力约 300～500 g/侧，牵引方向向前下约 15°～30°，'
     '每日佩戴面具约 12～14 小时（具体力值与时间依病例与装置而定）。', '加力：')
STEP(doc, '第 6 步', '复诊安排：加力型装置每 4～6 周复诊一次，'
     '评估磨牙移动、间隙变化、组织反应与口腔卫生，必要时更换牵引橡皮圈；'
     '单纯保持型每 3～6 个月复查。', '复诊：')
STEP(doc, '第 7 步', '影像学随访：每 6～12 个月拍摄影像，'
     '评估磨牙位置、牙根状况与恒牙萌出进程。', '随访：')
STEP(doc, '第 8 步', '疗程结束与拆除：达到治疗目标、恒牙萌出到位或需转入综合矫治时拆除装置，'
     '评估是否需要更换为保持装置或进入正畸治疗。', '拆除：')

NOTE(doc, '加力大小、方向与佩戴时间必须个体化设定，并记录于病历；'
          '力值过大可导致磨牙牙根吸收、牙髓反应、支抗丧失与装置脱落，'
          '力值不足则治疗周期延长、依从性下降。')

H2(doc, '5.6　粘接操作要点与常见失误')
B(doc, '隔湿是成功的关键：唾液污染是带环松动脱落最主要的可控制原因，'
       '建议使用开口器 + 吸唾器 + 棉卷，条件允许时使用橡皮障。')
B(doc, '带环必须在完全就位状态下粘固，边缘不得有悬突或压迫牙龈；'
       '粘接剂应充盈但不宜过厚，多余粘接剂须彻底清除。')
B(doc, '粘接后须再次检查咬合：带环或丝圈形成早接触可导致装置脱落、'
       '基牙不适与咬合创伤。')
B(doc, '避免丝圈或腭弓进入倒凹、压迫牙龈、跨越合面；'
       '否则将影响清洁并引起软组织炎症。')
B(doc, '腭托组织面必须与腭黏膜贴合且无高点，'
       '过高点可致压迫性溃疡、发音障碍与患儿拒戴。')

# ============================ 6 复诊维护 ============================
page_break(doc)
H1(doc, '6　复诊、维护与并发症处理')

H2(doc, '6.1　复诊计划')
B(doc, '戴入后 1 周：检查装置稳定性、组织反应、咬合与口腔卫生，强化清洁指导。')
B(doc, '戴入后 1 个月：评估适应情况、基牙与软组织状况。')
B(doc, '单纯保持型：此后每 3～6 个月复查一次，'
       '评估装置完整性、间隙维持效果、基牙与继承恒牙萌出进程。')
B(doc, '加力型（推磨牙、扩弓、前牵引）：每 4～6 周复诊，按设计加力或调整。')
B(doc, '每次复诊均须记录：装置状态、间隙测量值、软组织与菌斑情况、'
       '影像学变化及处理措施。')

H2(doc, '6.2　常见并发症与处理')
CAP(doc, '表 5　固定式间隙保持器与改良 Nance 弓的常见并发症及处理')
TABLE(doc,
      ['并发症', '常见原因', '处理原则'],
      [3.3, 5.5, 7.7],
      [['带环松动或脱落', '粘接面唾液污染、粘接剂不足、咬合力过大、'
        '基牙牙冠高度不足或形态不良',
        '清洁后重新粘固；反复脱落者考虑更换带环、改全冠或调整设计'],
       ['丝圈/舌弓变形、压迫牙龈或进入倒凹', '外力作用、制作不当、'
        '基牙或邻牙萌出导致相对位置改变',
        '口内调整或取下重做；解除对牙龈与软组织的压迫'],
       ['腭托压迫性溃疡、发音或吞咽不适', '腭托组织面不贴合、树脂过多、'
        '腭托过厚过大、恒牙萌出改变接触关系',
        '缓冲调磨组织面、缩小边缘；必要时重做腭托；'
        '持续不适者评估是否更换装置'],
       ['菌斑堆积、牙龈炎、带环下脱矿', '口腔卫生维护不足、'
        '装置形成滞留区、高龋活性',
        '洁治、局部用氟、强化菌斑控制；严重者拆除装置改活动式'],
       ['基牙牙髓症状或牙根吸收', '持续矫治力过大、'
        '粘接剂或带环对牙周的刺激、磨牙移动过快',
        '停止加力或拆除装置，行牙髓活力评估与相应治疗'],
       ['装置误吞或误吸（罕见）', '装置松动后未及时处理、'
        '患儿自行摆弄',
        '立即就医，行影像学检查并请相关科室会诊取出'],
       ['间隙保持失败（间隙已明显丧失）', '装置脱落未及时复诊、'
        '设计选择不当、缺隙评估失误',
        '重新评估，必要时正畸会诊，考虑序列拔牙或综合治疗'],
       ['横向或垂直向关系恶化', '对颌牙过度萌出、'
        '未同时控制咬合高度',
        '评估是否需增加咬合垫、平面导板或调整设计']],
      first_col_bold=True)

H2(doc, '6.3　患儿与家长宣教要点')
B(doc, '避免用装置侧咀嚼过硬、过黏食物（如硬糖、牛轧糖、口香糖），防止装置变形或脱落。')
B(doc, '每日早晚及餐后刷牙，重点清洁带环边缘、丝圈下方与腭托边缘；'
       '必要时使用单束刷、牙缝刷或冲牙器协助清洁。')
B(doc, '装置松动、变形、压迫牙龈或出现疼痛时，应及时就诊，'
       '不可自行调整或取下。')
B(doc, '按期复诊；复诊间隔过长是保持失败的重要原因。')
B(doc, '装置脱落并误吞时立即就医。')

# ============================ 7 拆除 ============================
H1(doc, '7　拆除指征与时机')
B(doc, '继承恒牙萌出至可与邻牙建立正常邻接关系，'
       '影像学显示牙根形成约 2/3、骨覆盖消失或临床可触及萌出牙冠。')
B(doc, '恒牙萌出至咬合平面并建立稳定咬合关系。')
B(doc, '需转入综合正畸治疗或序列拔牙计划，装置妨碍后续治疗。')
B(doc, '装置反复脱落、变形，或已引起不可逆的基牙/软组织损害。')
B(doc, '装置阻碍继承恒牙萌出、造成牙龈压迫或恒牙异位萌出。')
B(doc, '先天缺牙且正畸/修复方案已确定，需要重新分配或调整间隙。')

NOTE(doc, '拆除前应拍摄影像确认继承恒牙位置与萌出状态，'
          '避免过早拆除导致间隙再度丧失；拆除后仍需定期随访。')

# ============================ 8 质控 ============================
H1(doc, '8　质量控制要点')
B(doc, '适应证判断准确：确认继承恒牙距萌出时间充足，避免"过度保持"与"无效保持"。')
B(doc, '装置设计与基牙条件匹配：固位可靠、力线合理、组织避让充分。')
B(doc, '制作质量：焊点牢固、无锐边毛刺、抛光良好、腭托贴合。')
B(doc, '粘接质量：严格隔湿、完全就位、无悬突、咬合无干扰。')
B(doc, '随访依从：按计划复诊，及时处理松动、变形与组织并发症。')
B(doc, '病历记录完整：装置类型、材料规格、粘接日期、每次复诊发现与处理。')
B(doc, '多学科协作：涉及早期矫治目标时与正畸医师共同制定与评估方案。')

# ============================ 9 附录 ============================
page_break(doc)
H1(doc, '9　附录')

H2(doc, '附录 A　临床快速核对清单')
H3(doc, '术前')
B(doc, '□ 明确早失牙位、时间与原因；□ 评估龋活性与口腔卫生；'
       '□ 影像学确认恒牙发育阶段与萌出通道；□ 测量缺隙与牙弓长度；'
       '□ 评估基牙固位条件；□ 完成龋病与牙髓治疗；'
       '□ 知情同意与费用说明。')
H3(doc, '术中')
B(doc, '□ 带环完全就位、边缘密合无悬突；□ 印模中带环未移位；'
       '□ 焊点牢固、无变形；□ 丝圈/腭弓组织避让 0.5～1.0 mm；'
       '□ 腭托贴合、边缘圆钝；□ 咬合无干扰；□ 隔湿充分、粘接剂清除彻底。')
H3(doc, '术后')
B(doc, '□ 清洁方法示范并让家长复述；□ 饮食注意事项交代；'
       '□ 预约复诊时间；□ 病历记录完整；□ 必要时局部用氟。')

H2(doc, '附录 B　器械材料清单')
CAP(doc, '表 6　固定式间隙保持器与改良 Nance 弓的材料配置清单')
TABLE(doc,
      ['环节', '材料与器械'],
      [3.2, 13.3],
      [['临床检查与取模',
        '检查盘、口镜、探针、镊子、间隙测量尺、儿童托盘、藻酸盐印模材、'
        '调拌用具、超硬石膏、咬合记录蜡'],
       ['固位与粘接',
        '预成不锈钢带环（0.12～0.18 mm）、预成不锈钢全冠、带环推压器、扩大钳、'
        '玻璃离子/树脂加强玻璃离子水门汀、正畸粘接剂、酸蚀剂、隔湿与吸唾器材'],
       ['技工制作',
        '不锈钢丝（0.8/0.9/1.0 mm）、梯形钳、三喙钳、断丝钳、点焊机、'
        '银焊/微焊材料、焊媒、腭托树脂、蜡与雕刀、打磨抛光设备'],
       ['改良部件',
        '螺旋扩弓器、颊面管、牵引钩、推簧、镍钛丝、前方牵引面具与橡皮圈、'
        '（可选）腭部微种植钉'],
       ['防护与随访',
        '开口器、氟保护漆或含氟涂料、影像学检查（根尖片/曲面体层片）、'
        '复诊记录表与宣教资料']],
      first_col_bold=True)

H2(doc, '附录 C　家长告知要点（可直接用于沟通）')
B(doc, '为什么要戴：乳牙过早脱落会让旁边的牙齿倒向缺口，占掉恒牙的位置，'
       '以后恒牙可能长不出来或长歪；这个装置的作用就是"守住位置"。')
B(doc, '装上后要注意：不要咬硬糖、冰块、黏性食物；认真刷牙，'
       '特别是装置边缘和金属圈下面。')
B(doc, '可能发生什么：装置可能松动或脱落，牙龈可能被压到发红、'
       '口腔卫生不好时牙齿表面可能脱矿；按期复诊可以及早发现和处理。')
B(doc, '什么时候来：按预约时间复诊；如果装置松动、变形、'
       '压得牙龈疼或孩子说不舒服，请尽快来医院。')
B(doc, '什么时候取：等下面的恒牙长出来到合适的位置，'
       '或需要开始做牙齿矫正时，我们会把它取下来。')

# ============================ 参考文献 ============================
page_break(doc)
H1(doc, '参考文献与延伸阅读')
for i, r in enumerate([
    '中华口腔医学会儿童口腔医学专业委员会. 儿童间隙保持器的临床应用指南. 北京: 中华口腔医学会.',
    '葛立宏, 主编. 儿童口腔医学. 北京: 人民卫生出版社.',
    '赵志河, 主编. 口腔正畸学. 北京: 人民卫生出版社.',
    '秦满, 主编. 儿童口腔科临床操作教程. 北京: 人民卫生出版社.',
    'American Academy of Pediatric Dentistry. Management of the Developing Dentition '
    'and Occlusion in Pediatric Dentistry. The Reference Manual of Pediatric Dentistry. '
    'Chicago: AAPD.',
    'American Academy of Pediatric Dentistry. Policy on Management of the Developing '
    'Dentition and Occlusion in Pediatric Dentistry. Chicago: AAPD.',
    'Proffit WR, Fields HW, Larson BE, Sarver DM. Contemporary Orthodontics. '
    'Elsevier.',
    'Ngan P, Alkire RG, Fields H. Management of space problems in the primary and '
    'mixed dentitions. Journal of the American Dental Association.',
    'Kupietzky A, Tal E. The transpalatal arch: an alternative to the Nance appliance '
    'for space maintenance. Pediatric Dentistry.',
    'Bhattacharjee B, et al. A review on space maintainers in pediatric dentistry. '
    'Journal of International Oral Health.',
], 1):
    p = doc.add_paragraph(style='CN-Ref')
    pf = p.paragraph_format
    pf.left_indent = Cm(1.25)
    pf.first_line_indent = Cm(-0.75)
    set_run(p.add_run('[%d] ' % i), size=9.5, bold=True)
    set_run(p.add_run(r), size=9.5)

NOTE(doc, '文献著录仅列示常用指南与教材，未逐一标注版本、年份与页码；'
          '临床决策请查阅最新版指南原文与教材。')

doc.save('固定式间隙保持器与改良Nance弓矫治器_临床规范参考.docx')
print('DOCX saved')
