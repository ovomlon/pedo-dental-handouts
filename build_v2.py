# -*- coding: utf-8 -*-
"""图解版临床手册：以示意图为主、文字精简、附病例举例。生成 DOCX。"""
import os
import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_EA = '宋体'
HEAD_EA = '微软雅黑'
BODY_ASCII = 'Times New Roman'
HEAD_ASCII = 'Arial'
DARK = RGBColor(0x1F, 0x36, 0x4D)
ACCENT = RGBColor(0x0F, 0x5C, 0x8A)
GREY = RGBColor(0x5F, 0x67, 0x70)
RED = RGBColor(0xB0, 0x3A, 0x2B)
FIGS = 'figs'
TODAY = datetime.date.today().strftime('%Y-%m-%d')


def set_run(run, size=10.5, ea=BODY_EA, ascii_f=BODY_ASCII, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_f
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rFonts.set(qn(k), ascii_f)
    rFonts.set(qn('w:eastAsia'), ea)
    return run


def style_font(style, size, ea, ascii_f, bold=False, color=None):
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.name = ascii_f
    if color is not None:
        style.font.color.rgb = color
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rFonts.set(qn(k), ascii_f)
    rFonts.set(qn('w:eastAsia'), ea)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def cell_margins(cell, top=50, bottom=50, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('start', left), ('bottom', bottom), ('end', right)):
        node = OxmlElement('w:' + tag)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tcPr.append(mar)


def add_field(paragraph, instr, size=9):
    run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    set_run(run, size=size)


def make_styles(doc):
    st = doc.styles
    style_font(st['Normal'], 10.5, BODY_EA, BODY_ASCII)

    def mk(name, size, ea, ascii_f, bold, color, before, after, line=1.35,
           left=0.0, indent_chars=0, keep_next=False):
        s = st.add_style(name, 1)
        s.base_style = st['Normal']
        style_font(s, size, ea, ascii_f, bold, color)
        pf = s.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line
        if left:
            pf.left_indent = Cm(left)
        if indent_chars:
            pPr = s.element.get_or_add_pPr()
            ind = pPr.find(qn('w:ind'))
            if ind is None:
                ind = OxmlElement('w:ind')
                pPr.insert_element_before(
                    ind, 'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap',
                    'w:jc', 'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap',
                    'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr',
                    'w:pPrChange')
            ind.set(qn('w:firstLineChars'), str(int(indent_chars * 100)))
        if keep_next:
            s.paragraph_format.keep_with_next = True
        return s

    mk('T-Title', 19, HEAD_EA, HEAD_ASCII, True, DARK, 0, 6, 1.15)
    mk('T-Sub', 12, HEAD_EA, HEAD_ASCII, False, ACCENT, 0, 4, 1.2)
    mk('T-H1', 14.5, HEAD_EA, HEAD_ASCII, True, DARK, 14, 6, 1.2, keep_next=True)
    mk('T-H2', 11.5, HEAD_EA, HEAD_ASCII, True, ACCENT, 9, 4, 1.2, keep_next=True)
    mk('T-Body', 10.5, BODY_EA, BODY_ASCII, False, None, 0, 4, 1.35)
    mk('T-Bul', 10.5, BODY_EA, BODY_ASCII, False, None, 0, 2, 1.3, left=0.7)
    mk('T-Cap', 9, HEAD_EA, HEAD_ASCII, False, GREY, 2, 8, 1.2)
    mk('T-TblHead', 9, HEAD_EA, HEAD_ASCII, True, RGBColor(0xFF, 0xFF, 0xFF), 0, 0, 1.15)
    mk('T-TblCell', 9, BODY_EA, BODY_ASCII, False, None, 0, 0, 1.18)
    mk('T-Note', 9.5, BODY_EA, BODY_ASCII, False, GREY, 3, 5, 1.3, left=0.4)


doc = Document()
make_styles(doc)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.header_distance = Cm(1.1)
sec.footer_distance = Cm(0.9)

hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(hp.add_run('固定式间隙保持器与改良 Nance 弓矫治器 · 图解临床手册'), size=8.5,
        ea=HEAD_EA, ascii_f=HEAD_ASCII, color=GREY)
pPr = hp._p.get_or_add_pPr()
pbdr = OxmlElement('w:pBdr')
bot = OxmlElement('w:bottom')
for k, v in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '2'),
             ('w:color', 'BFBFBF')):
    bot.set(qn(k), v)
pbdr.append(bot)
pPr.insert_element_before(pbdr, 'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku',
                          'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct',
                          'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi',
                          'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
                          'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap',
                          'w:jc', 'w:textDirection', 'w:textAlignment',
                          'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle',
                          'w:rPr', 'w:sectPr', 'w:pPrChange')
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(fp.add_run('第 '), size=9)
add_field(fp, 'PAGE')
set_run(fp.add_run(' 页 / 共 '), size=9)
add_field(fp, 'NUMPAGES')
set_run(fp.add_run(' 页　|　图解手册 · 临床参考'), size=8.5, color=GREY)


def H1(t):
    p = doc.add_paragraph(style='T-H1')
    set_run(p.add_run(t))
    return p


def H2(t):
    p = doc.add_paragraph(style='T-H2')
    set_run(p.add_run(t))
    return p


def P(t, size=10.5, color=None, bold=False):
    p = doc.add_paragraph(style='T-Body')
    set_run(p.add_run(t), size=size, color=color, bold=bold)
    return p


def BUL(t):
    p = doc.add_paragraph(style='T-Bul')
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    set_run(p.add_run('· ' + t))
    return p


def IMG(path, width_cm=16.6, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(os.path.join(FIGS, path), width=Cm(width_cm))
    if caption:
        c = doc.add_paragraph(style='T-Cap')
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(c.add_run(caption), size=9, ea=HEAD_EA, ascii_f=HEAD_ASCII, color=GREY)


def TABLE(cols, widths, rows, first_bold=False, size=9):
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, c in enumerate(cols):
        cell = t.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(c), size=size, ea=HEAD_EA, ascii_f=HEAD_ASCII, bold=True,
                color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, '1F364D')
        cell_margins(cell)
    for r in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(r):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            bold = first_bold and i == 0
            col = None
            if txt.startswith('√'):
                col = RGBColor(0x18, 0x64, 0x3E)
            elif txt.startswith('×'):
                col = RED
            set_run(p.add_run(txt), size=size, bold=bold, color=col)
            cell_margins(cells[i])
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ============================== 封面 ==============================
p = doc.add_paragraph(style='T-Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(p.add_run('固定式间隙保持器与改良 Nance 弓矫治器'), size=19, ea=HEAD_EA,
        ascii_f=HEAD_ASCII, bold=True, color=DARK)
p = doc.add_paragraph(style='T-Sub')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(p.add_run('图解临床手册　·　适应证 / 器械材料 / 操作步骤 / 禁忌证 / 病例举例'),
        size=12, ea=HEAD_EA, ascii_f=HEAD_ASCII, color=ACCENT)
p = doc.add_paragraph(style='T-Cap')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(p.add_run('适用于儿童口腔科 · 口腔全科 · 早期矫治　|　版本 %s' % TODAY),
        size=9, ea=HEAD_EA, ascii_f=HEAD_ASCII, color=GREY)

IMG('fig01_overview.png', 16.6, '图 1　固定式间隙保持器主要类型总览（①—⑥）')

TABLE(['类型', '一句话适应证', '支抗牙'],
      [3.4, 9.6, 3.6],
      [['① 带环丝圈式', '单侧 1 个牙缺失，恒牙 6 个月内不能萌出', '缺隙远中邻牙'],
       ['② 全冠丝圈式', '基牙大面积缺损或已做全冠', '全冠修复的基牙'],
       ['③ 远端靴式', '第一恒磨牙未萌时第二乳磨牙早失', '第一乳磨牙'],
       ['④ 固定舌弓', '双侧下颌乳磨牙早失', '下颌第一恒磨牙'],
       ['⑤ Nance 腭弓', '双侧上颌乳磨牙早失', '上颌第一恒磨牙'],
       ['⑥ 改良 Nance 弓', '保持间隙 + 早期矫治（推磨牙 / 前牵引 / 扩弓）',
        '上颌第一恒磨牙（可加骨支抗）']],
      first_bold=True)

# ============================== 1 结构与尺寸 ==============================
# 连续排版（避免大面积空白）
H1('1　结构与关键尺寸')
IMG('fig02_band_loop.png', 16.6, '图 2　带环丝圈式：结构要点与关键尺寸（颊面观）')
IMG('fig03_lingual_arch.png', 16.0, '图 3　固定舌弓：结构要点（下颌 · 合面观）')

# 连续排版（避免大面积空白）
IMG('fig04_nance.png', 16.6, '图 4　Nance 腭弓：合面观与腭部矢状剖面')

H2('常用尺寸速查')
TABLE(['部位 / 环节', '常用范围', '要点'],
      [4.2, 4.6, 7.8],
      [['带环厚度', '0.12–0.18 mm', '边缘密合、无悬突'],
       ['丝圈 / 舌弓钢丝', '0.8–0.9 mm', '刚性足够、便于调整'],
       ['腭弓钢丝', '0.9–1.0 mm', '贴合腭穹窿走行'],
       ['腭托厚度', '2–3 mm', '边缘圆钝、组织面均匀贴合'],
       ['装置距牙龈 / 黏膜', '0.5–1.0 mm', '不压迫、不进入倒凹'],
       ['推磨牙加力', '0.5–1.0 mm / 次', '每 4–6 周复诊'],
       ['前方牵引', '300–500 g/侧，向前下 15°–30°', '12–14 h/日'],
       ['保持型复查间隔', '3–6 个月', '间隙测量 + 装置检查']],
      first_bold=True)

# ============================== 2 改良 Nance 弓 ==============================
# 连续排版（避免大面积空白）
H1('2　改良 Nance 弓矫治器')
P('在经典 Nance 腭弓（上颌第一恒磨牙带环 + 腭弓 + 前腭托）基础上增加功能部件，'
  '使“保持间隙”与“早期矫治”合一。', size=10.5)
IMG('fig05_modified_nance.png', 16.6, '图 5　改良 Nance 弓的四种常见改良形式')

TABLE(['改良形式', '结构变化', '主要目的'],
      [3.6, 6.4, 6.6],
      [['① 推磨牙向远中', '腭弓加圈曲 / 推簧，或焊远中向牵引臂',
        '上颌第一恒磨牙异位萌出、早期 II 类'],
       ['② 前方牵引支抗', '带环焊颊面管 / 牵引钩，腭托加大',
        '安氏 III 类上颌后缩的腭部支抗'],
       ['③ 扩弓—保持', '腭中缝加螺旋扩弓器', '横向不足合并间隙丧失'],
       ['④ 骨支抗 / 固位改良', '腭部微种植钉；基牙改全冠、加邻牙支托',
        '需绝对支抗或基牙条件差']],
      first_bold=True)

# ============================== 3 怎么选 ==============================
# 连续排版（避免大面积空白）
H1('3　装置怎么选')
IMG('fig06_decision_flow.png', 15.6, '图 6　装置选择决策流程')

TABLE(['临床情形', '首选', '关键点'],
      [7.0, 4.2, 5.4],
      [['单侧下颌第一乳磨牙早失，第二乳磨牙完好', '带环丝圈式', '带环于第二乳磨牙'],
       ['单侧第二乳磨牙早失，第一恒磨牙未萌', '远端靴式', '引导恒磨牙萌出，时机要求高'],
       ['双侧下颌乳磨牙早失', '固定舌弓', '支抗牙 36 / 46'],
       ['双侧上颌乳磨牙早失', 'Nance 腭弓', '腭托需均匀贴合'],
       ['上颌第一恒磨牙异位萌出', '改良 Nance 弓 + 推簧', '尽早干预，评估可否复位'],
       ['上颌后缩拟行前方牵引', '改良 Nance 弓 + 面具', '需保证每日佩戴时间'],
       ['上颌横向不足合并缺隙', '改良 Nance 弓 + 扩弓器', '扩弓与保持同步'],
       ['口腔卫生极差、龋病未控制', '先治龋 / 暂缓固定装置', '必要时改活动式']],
      first_bold=True)

# ============================== 4 操作步骤 ==============================
# 连续排版（避免大面积空白）
H1('4　操作步骤与器械材料')
IMG('fig07_steps.png', 16.6, '图 7　操作流程与关键点（固定式保持器 / 改良 Nance 弓通用）')

TABLE(['步骤', '操作要点'],
      [3.4, 13.2],
      [['1 术前评估', '影像判断恒牙发育阶段与萌出通道；测缺隙；查基牙与口腔卫生'],
       ['2 试带环', '就位、边缘密合无悬突、邻接恢复、咬合无干扰'],
       ['3 取模灌模', '带环随印模取出并牢固复位；超硬石膏灌制'],
       ['4 技工制作', '丝圈 / 舌弓 / 腭弓弯制 → 焊接带环 → 制作腭托 → 抛光'],
       ['5 试戴调整', '丝圈与腭弓距组织 0.5–1.0 mm；腭托均匀贴合；咬合无干扰'],
       ['6 粘接', '严格隔湿；完全就位；去净多余粘接剂；复查咬合'],
       ['7 医嘱宣教', '避免硬黏食物；重点清洁带环边缘、丝圈下方、腭托边缘'],
       ['8 定期复查', '保持型 3–6 个月；加力型 4–6 周；记录间隙测量值']],
      first_bold=True)

H2('器械与材料清单')
TABLE(['环节', '材料与器械'],
      [3.4, 13.2],
      [['检查与取模', '检查盘、间隙测量尺、儿童托盘、藻酸盐印模材、超硬石膏、咬合蜡'],
       ['固位与粘接', '预成不锈钢带环、预成不锈钢全冠、玻璃离子 / 树脂加强玻璃离子、'
        '正畸粘接剂、开口器与吸唾器'],
       ['技工制作', '不锈钢丝 0.8 / 0.9 / 1.0 mm、梯形钳、三喙钳、断丝钳、点焊机、'
        '银焊与焊媒、腭托树脂、打磨抛光器材'],
       ['改良部件', '螺旋扩弓器、颊面管、牵引钩、推簧、镍钛丝、前方牵引面具与橡皮圈、'
        '（可选）腭部微种植钉'],
       ['防护与随访', '氟保护漆 / 含氟涂料、根尖片与曲面体层片、复诊记录表、宣教资料']],
      first_bold=True)

# ============================== 5 适应证 / 禁忌证 ==============================
doc.add_page_break()
H1('5　适应证与禁忌证')
TABLE(['√ 适应证', '× 禁忌证 / 慎用'],
      [8.3, 8.3],
      [['√ 乳磨牙早失，恒牙 6 个月内不能萌出', '× 继承恒牙即将萌出（牙根 > 2/3）'],
       ['√ 单侧 1 牙缺失（带环 / 全冠丝圈式）', '× 龋病、牙髓病未控制'],
       ['√ 双侧缺隙（舌弓 / Nance 腭弓）', '× 口腔卫生极差、高龋活性'],
       ['√ 第一恒磨牙未萌 + 第二乳磨牙早失（远端靴式）', '× 基牙牙冠高度 / 形态不足'],
       ['√ 先天缺牙需维持间隙', '× 腭穹窿过浅、腭部瘢痕或感觉障碍'],
       ['√ 依从性差，不宜活动式装置', '× 对镍铬 / 不锈钢过敏'],
       ['√ 需加强支抗（推磨牙 / 前牵引 / 扩弓）', '× 无法定期复诊'],
       ['√ 磨牙异位萌出的早期阻断', '× 需综合正畸 / 拔牙矫治'],
       ['√ 腭部阻挡以纠正不良习惯', '× 面具佩戴时间无法保证（改良 Nance 弓）']],
      size=9)

IMG('fig08_dos_donts.png', 16.6, '图 8　常见错误对照（左列正确 · 右列错误）')

# ============================== 6 复诊与并发症 ==============================
# 连续排版（避免大面积空白）
H1('6　复诊、维护与并发症')
IMG('fig09_timeline.png', 16.6, '图 9　复诊与维护时间轴')

TABLE(['并发症', '常见原因', '处理'],
      [3.4, 6.2, 7.0],
      [['带环松动 / 脱落', '隔湿不良、咬合力过大、基牙条件差', '重新粘固；反复脱落改全冠或换设计'],
       ['丝圈变形、压龈', '外力、制作不当、邻牙萌出', '口内调整或重做，解除压迫'],
       ['腭托压迫溃疡', '组织面不贴合、树脂过多', '缓冲调磨；必要时重做腭托'],
       ['牙龈炎、带环下脱矿', '卫生维护差、滞留区', '洁治 + 用氟 + 强化宣教；必要时拆除'],
       ['牙髓症状 / 牙根吸收', '矫治力过大、移动过快', '停止加力或拆除，评估牙髓'],
       ['间隙保持失败', '脱落未复诊、设计不当', '重新评估，必要时正畸会诊']],
      first_bold=True)

H2('家长沟通要点（可直接照读）')
BUL('这个装置是“占位”的：乳牙早掉，旁边的牙会倒过来占掉恒牙的位置。')
BUL('不要咬硬糖、冰块、黏性食物；刷牙要重点刷装置边缘和金属圈下面。')
BUL('装置松动、变形、牙龈被压红或孩子喊疼，请尽快来医院，不要自己掰。')
BUL('按预约复查；恒牙长到位置或开始做矫正时就把它取下来。')

# ============================== 7 拆除时机 ==============================
# 连续排版（避免大面积空白）
H1('7　拆除时机')
IMG('fig11_removal.png', 16.0, '图 10　拆除时机判断（影像学 + 临床）')

# ============================== 8 病例举例 ==============================
H1('8　病例举例')
IMG('fig12_cases.png', 16.6, '图 11　四个典型病例示意（红圈 = 缺隙 / 异位萌出）')

TABLE(['病例', '年龄 / 情况', '影像学与检查', '装置选择', '理由与要点'],
      [1.5, 3.6, 3.7, 3.6, 4.2],
      [['1', '6 岁 4 个月，男；74 因龋拔除', '75 完好；继承恒牙距萌出 > 1 年',
        '带环丝圈式（带环于 75）', '单侧单牙缺失；丝圈抵 73，避免 75 近中倾斜'],
       ['2', '7 岁 1 个月，女；55、65 早失', '腭穹窿较深；16、26 已萌',
        'Nance 腭弓', '跨中线同时保持双侧；腭托提供支抗'],
       ['3', '7 岁 8 个月，男；16 异位萌出', '可复性嵌顿，牙根形成约 1/2',
        '改良 Nance 弓 + 推簧', '早期阻断；避免 55 早脱后间隙进一步丧失'],
       ['4', '8 岁 2 个月，女；上颌后缩、反合', 'ANB < 0°，上颌后缩',
        '改良 Nance 弓 + 颊面管 + 面具', '腭部支抗 + 前牵引 300–500 g/侧'],
       ['5', '5 岁 10 个月，女；85 早失', '46 尚未萌出，骨覆盖完整',
        '远端靴式', '引导 46 在正常位置萌出；拔牙前取模'],
       ['6', '6 岁 8 个月，男；75、85 早失', '36、46 已萌，牙冠条件好',
        '固定舌弓', '下颌多牙缺失，保持牙弓长度与磨牙宽度'],
       ['7', '9 岁，女；上颌横向不足 + 55 早失', '后牙反合，上颌宽度不足',
        '改良 Nance 弓 + 螺旋扩弓器', '扩弓与间隙保持同步，注意加力节奏'],
       ['8', '8 岁 6 个月，男；75 早失，36 已萌', '36 完全萌出，牙冠高度足够',
        '带环丝圈式（带环于 36）', '恒磨牙已萌，固位可靠，制作简单'],
       ['9', '7 岁 3 个月，男；74 早失但多牙龋坏', '口腔卫生极差，多牙需治疗',
        '先治龋 + 口腔卫生指导，暂缓固定装置', '卫生与依从性不达标；必要时改活动式'],
       ['10', '6 岁 2 个月，女；65 早失 + 腭部瘢痕', '腭穹窿浅，腭部术后瘢痕',
        '改用带环丝圈式（单侧）', 'Nance 腭托易压迫、脱位，应回避']],
      first_bold=True, size=8.5)

# ============================== 9 附录 ==============================
# 连续排版（避免大面积空白）
H1('9　附录')
H2('术前 / 术中 / 术后核对清单')
BUL('术前：影像确认恒牙发育阶段；测量缺隙；评估基牙与菌斑；完成龋病治疗；知情同意。')
BUL('术中：带环完全就位无悬突；焊点牢固；丝圈 / 腭弓距组织 0.5–1.0 mm；'
    '腭托贴合；咬合无干扰；隔湿充分。')
BUL('术后：清洁方法示范并让家长复述；交代饮食禁忌；预约复查；病历记录材料规格。')

H2('主要参考依据')
for i, r in enumerate([
    '中华口腔医学会儿童口腔医学专业委员会. 儿童间隙保持器的临床应用指南.',
    '葛立宏, 主编. 儿童口腔医学. 人民卫生出版社.',
    '赵志河, 主编. 口腔正畸学. 人民卫生出版社.',
    '秦满, 主编. 儿童口腔科临床操作教程. 人民卫生出版社.',
    'American Academy of Pediatric Dentistry. Management of the Developing Dentition '
    'and Occlusion in Pediatric Dentistry. The AAPD Reference Manual.',
    'Proffit WR, et al. Contemporary Orthodontics. Elsevier.',
], 1):
    p = doc.add_paragraph(style='T-Note')
    p.paragraph_format.left_indent = Cm(1.15)
    p.paragraph_format.first_line_indent = Cm(-0.7)
    set_run(p.add_run('[%d] ' % i), size=9, bold=True)
    set_run(p.add_run(r), size=9)

p = doc.add_paragraph(style='T-Note')
set_run(p.add_run('说明：图中尺寸、力值与复诊间隔均为临床常用参考范围，'
                  '须结合患者年龄、牙列阶段、影像学表现与依从性个体化调整；'
                  '具体产品以说明书与注册证为准。'), size=9, color=GREY)

doc.save('固定式间隙保持器与改良Nance弓_图解手册.docx')
print('OK saved')
